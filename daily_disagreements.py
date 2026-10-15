"""Race-day sheet: champion vs SAS-tree challenger AFTER scratches, with union legs and disagreements.
Scratches/conditions come from a saved Equibase RSS snapshot (feed_archive/<date>/*.rss) parsed with the
production scratches.py logic, or from --scratches "1:1;2:6,7" if no feed is available.
Usage: python daily_disagreements.py --date 20261014 [--rss feed_archive/20261014/KEE_manual.rss] [--scratches "..."] [--cond "dirt=Fast,turf=Firm"]"""
import sys, re, glob, argparse, warnings; warnings.simplefilter("ignore")
from pathlib import Path
import numpy as np, pandas as pd
HERE = Path(__file__).resolve().parent; sys.path.insert(0, str(HERE))
CH = HERE.parent / "Challenger" / "data"
ap = argparse.ArgumentParser(); ap.add_argument("--date", required=True); ap.add_argument("--track", default="KEE")
ap.add_argument("--rss"); ap.add_argument("--scratches", default=""); ap.add_argument("--cond", default="")
ap.add_argument("--gap", type=float, default=6.0, help="pts of P(win) that counts as a disagreement")
a = ap.parse_args()
d = a.date; mm, yyyy = d[4:6], d[:4]

# ---- scratches + conditions ----
scr = {}; cond = {}; offturf = set(); rss_note = ""
rss = a.rss or (sorted(glob.glob(str(HERE / "feed_archive" / d / f"{a.track}_*.rss"))) or [None])[-1]
if rss:
    import scratches as S
    raw = Path(rss).read_bytes()
    S._fetch_raw_feed = lambda url, timeout: raw          # parse the snapshot through the production parser
    for e in S.get_scratches(a.track, d[4:], yyyy):
        scr.setdefault(int(e.race), set()).add(str(e.program_number).strip())
    st = S.get_rss_status(a.track); cond = {k: v for k, v in st.items() if k in ("dirt", "turf", "aw") and v}
    for c in S.fetch_all_changes(a.track):
        t = (getattr(c, "raw_text", "") or "")
        if re.search(r"off the turf|off-turf|turf changed to dirt|moved to (the )?(main|dirt)", t, re.I) and c.race: offturf.add(int(c.race))
    rss_note = f"feed snapshot {Path(rss).name}, last bulletin {st.get('last_updated_et')}"
for part in filter(None, a.scratches.split(";")):
    r, pns = part.split(":"); scr.setdefault(int(r), set()).update(p.strip() for p in pns.split(","))
for part in filter(None, a.cond.split(",")):
    k, v = part.split("="); cond[k.strip()] = v.strip()

# ---- score both models on the post-scratch field ----
import config, model_setup, score, model_registry as mr
from ingest_drf import load_drf; from features import engineer_features
import lightgbm as lgb
model_setup.setup_registry(config)
cdir = Path(config.COEFF_DIR); cdir = cdir if cdir.exists() else HERE / "coefficients"
drf = load_drf(str(HERE.parent / a.track / "RAW_DATA" / "RACINGFORM" / yyyy / f"{a.track}{mm}{d[6:]}.DRF"), a.track, d[4:], yyyy, validate=False)
drf["pn"] = drf["ProgramNumberifavailable"].astype(str).str.strip(); drf["race"] = pd.to_numeric(drf["Race"], errors="coerce").astype(int)
pre = len(drf)
keep = ~drf.apply(lambda r: r["pn"] in scr.get(r["race"], set()), axis=1)
drf = drf[keep].reset_index(drop=True)
feats = engineer_features(drf.copy())
sc = mr.get_scoring_models(a.track, config, d)
cp = score.run_scoring(feats.copy(), cdir, sc)
cp["pn"] = cp["ProgramNumberifavailable"].astype(str).str.strip(); cp["race"] = cp["Race"].astype(int)
models = {s: lgb.Booster(model_file=str(CH / f"sastree_{s}.txt")) for s in ["dirt", "turf", "maiden"]}
feat = models["dirt"].feature_name()
x = feats.copy(); x["pn"] = x["ProgramNumberifavailable"].astype(str).str.strip(); x["race"] = pd.to_numeric(x["Race"], errors="coerce").astype(int)
for v in feat: x[v] = pd.to_numeric(x[v], errors="coerce") if v in x.columns else np.nan
S_ = x["Surface"].astype(str); RT = x["RaceType"].astype(str)
x["seg"] = np.where(RT.isin(["M", "S"]), "maiden", np.where(S_.isin(["T", "t"]), "turf", "dirt")); x["raw"] = np.nan
for s, m in models.items():
    k = x["seg"] == s
    if k.any(): x.loc[k, "raw"] = m.predict(x.loc[k, feat].values.astype(np.float32))
x["p_ch"] = x["raw"] / x.groupby("race")["raw"].transform("sum")
df = cp[["race", "pn", "HorseName", "Surface", "RaceType", "MornLineOddsifavailable", "MornOddsAdj", "ProbToWin"]].rename(columns={"ProbToWin": "p_cp"}).merge(x[["race", "pn", "p_ch"]], on=["race", "pn"])
for p in ["p_cp", "p_ch"]:
    s_ = df.sort_values(["race", p], ascending=[True, False])
    df["rank_" + p] = s_.groupby("race").cumcount().add(1).reindex(df.index)
    df["pa_" + p] = (s_.groupby("race")[p].cumsum() - s_[p]).reindex(df.index)
    df["gold_" + p] = ((pd.to_numeric(df["MornOddsAdj"], errors="coerce") + 1) * df[p] >= 1.5) & (df["pa_" + p] < 0.5)
    df["top_" + p] = df["pa_" + p] < 0.5
df["dp"] = (df["p_ch"] - df["p_cp"]) * 100
df["odds_cp"] = 1 / df["p_cp"] - 1; df["odds_ch"] = 1 / df["p_ch"] - 1
out = HERE / "results" / f"daily_{a.track}_{d}.csv"; df.sort_values(["race", "p_cp"], ascending=[True, False]).to_csv(out, index=False)

# ---- report ----
L = []
L.append(f"{a.track} {d} — post-scratch sheet: champion vs tree   [{rss_note or 'manual scratches'}]")
L.append(f"conditions: " + (", ".join(f"{k} {v}" for k, v in cond.items()) if cond else "not posted yet") + (f"   OFF THE TURF: R{sorted(offturf)}" if offturf else ""))
_pk = lambda z: int(re.sub(r"\D", "", z) or 0)
L.append(f"scratches ({pre - len(drf)}): " + "; ".join("R%d: #%s" % (r, ",".join(sorted(v, key=_pk))) for r, v in sorted(scr.items())))
L.append("")
L.append("legend: ★ = both models' #1 (single candidate) · U = on the union ticket (top-50% mass of either model) · G/g = gold champion/tree · Δ = tree − champion, pts")
for r, g in df.groupby("race"):
    g = g.sort_values("p_cp", ascending=False)
    surf = g["Surface"].iloc[0]; rt = g["RaceType"].iloc[0]
    c = cond.get("turf" if surf == "T" else "dirt")
    flag = "  ⚠ OFF TURF — out of scope" if r in offturf else ("" if not c or c.title() in ("Fast", "Firm", "Good") else f"  ⚠ {c} — off-track, not modeled")
    t1 = g.sort_values("p_cp").iloc[-1]["pn"]; t2 = g.sort_values("p_ch").iloc[-1]["pn"]
    union = g[g["top_p_cp"] | g["top_p_ch"]]
    L.append(f"\nR{r} {surf}/{rt} ({len(g)} run){flag}")
    L.append(f"  union ticket: {' '.join('#'+p for p in union.sort_values('p_cp', ascending=False)['pn'])}" + (f"   ★ single candidate #{t1} (both #1)" if t1 == t2 else f"   #1 split: champ #{t1} / tree #{t2}"))
    L.append(f"  {'pn':>3s} {'horse':20s} {'ML':>5s} | {'champ':>6s} {'odds':>5s} rk    | {'tree':>6s} {'odds':>5s} rk    | {'Δ':>6s}")
    for _, h in g.iterrows():
        fc = ("G" if h["gold_p_cp"] else ("U" if h["top_p_cp"] else " ")); ft = ("g" if h["gold_p_ch"] else ("U" if h["top_p_ch"] else " "))
        mark = "  <<" if abs(h["dp"]) >= a.gap else ""
        L.append(f"  {h['pn']:>3s} {str(h['HorseName'])[:20]:20s} {str(h['MornLineOddsifavailable']):>5s} | {h['p_cp']*100:5.1f}% {h['odds_cp']:5.1f} {int(h['rank_p_cp']):2d} {fc}  | {h['p_ch']*100:5.1f}% {h['odds_ch']:5.1f} {int(h['rank_p_ch']):2d} {ft}  | {h['dp']:+6.1f}{mark}")
big = df[df["dp"].abs() >= a.gap].sort_values("dp", key=abs, ascending=False)
L.append("\n== Disagreements (|Δ| ≥ %.0f pts), biggest first ==" % a.gap)
for _, h in big.iterrows():
    side = "TREE likes" if h["dp"] > 0 else "CHAMP likes"
    L.append(f"  R{int(h['race'])} #{h['pn']:>2s} {str(h['HorseName'])[:20]:20s} ML {str(h['MornLineOddsifavailable']):>4s}  champ {h['p_cp']*100:4.1f}% (rk {int(h['rank_p_cp'])})  tree {h['p_ch']*100:4.1f}% (rk {int(h['rank_p_ch'])})  {side} {abs(h['dp']):.0f} pts")
txt = "\n".join(L); print(txt)
(HERE / "results" / f"daily_{a.track}_{d}.txt").write_text(txt)

# ---- Excel ----
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
wb = Workbook(); ws = wb.active; ws.title = f"{a.track} {d[4:6]}-{d[6:]}"
hdr = Font(bold=True); thin = Side(style="thin", color="999999")
fill_cp = PatternFill("solid", fgColor="FFF2CC"); fill_ch = PatternFill("solid", fgColor="DDEBF7"); fill_both = PatternFill("solid", fgColor="E2EFDA")
fill_race = PatternFill("solid", fgColor="D9D9D9"); fill_dis = PatternFill("solid", fgColor="FCE4D6")
ws.append([f"{a.track} {d} — post-scratch: champion vs tree", "", "", "", f"conditions: " + (", ".join(f"{k} {v}" for k, v in cond.items()) if cond else "not posted")]); ws["A1"].font = Font(bold=True, size=12)
ws.append(["scratches: " + "; ".join("R%d: #%s" % (r, ",".join(sorted(v, key=_pk))) for r, v in sorted(scr.items()))])
ws.append(["Yellow = champion gold · Blue = tree gold · Green = gold by both · Orange Δ = disagreement ≥ %d pts · U = on union ticket · ★ = both models' #1" % a.gap])
ws.append([])
cols = ["Race", "PN", "Horse", "ML", "ML adj", "Champ P(win)", "Champ odds", "Champ rk", "Champ", "Tree P(win)", "Tree odds", "Tree rk", "Tree", "Δ pts", "Union", "Single ★"]
ws.append(cols); hr = ws.max_row
for c in ws[hr]: c.font = hdr; c.alignment = Alignment(horizontal="center", wrap_text=True); c.border = Border(bottom=thin)
for r, g in df.groupby("race"):
    g = g.sort_values("p_cp", ascending=False)
    t1 = g.sort_values("p_cp").iloc[-1]["pn"]; t2 = g.sort_values("p_ch").iloc[-1]["pn"]
    c_ = cond.get("turf" if g["Surface"].iloc[0] == "T" else "dirt")
    note = "OFF TURF — out of scope" if r in offturf else ("" if not c_ or c_.title() in ("Fast", "Firm", "Good") else f"{c_} — off-track, not modeled")
    union = g[g["top_p_cp"] | g["top_p_ch"]].sort_values("p_cp", ascending=False)["pn"].tolist()
    ws.append([f"R{r}  {g['Surface'].iloc[0]}/{g['RaceType'].iloc[0]}  {len(g)} run", "", f"union: " + " ".join("#" + p for p in union), "", ("★ single #" + t1) if t1 == t2 else f"#1 split: champ #{t1} / tree #{t2}", "", "", "", "", "", "", "", "", "", "", note])
    for c in ws[ws.max_row]: c.fill = fill_race; c.font = Font(bold=True)
    for _, h in g.iterrows():
        ws.append([r, h["pn"], h["HorseName"], h["MornLineOddsifavailable"], h["MornOddsAdj"], h["p_cp"], h["odds_cp"], int(h["rank_p_cp"]),
                   "G" if h["gold_p_cp"] else ("U" if h["top_p_cp"] else ""), h["p_ch"], h["odds_ch"], int(h["rank_p_ch"]),
                   "G" if h["gold_p_ch"] else ("U" if h["top_p_ch"] else ""), h["dp"], "U" if (h["top_p_cp"] or h["top_p_ch"]) else "", "★" if (t1 == t2 and h["pn"] == t1) else ""])
        row = ws[ws.max_row]
        f = fill_both if (h["gold_p_cp"] and h["gold_p_ch"]) else (fill_cp if h["gold_p_cp"] else (fill_ch if h["gold_p_ch"] else None))
        if f:
            for c in row: c.fill = f
        if abs(h["dp"]) >= a.gap: row[13].fill = fill_dis; row[13].font = Font(bold=True)
        row[5].number_format = "0.0%"; row[9].number_format = "0.0%"
        for j in (4, 6, 10): row[j].number_format = "0.0"
        row[13].number_format = "+0.0;-0.0;0.0"
for i, w in enumerate([6, 4, 22, 6, 7, 11, 10, 8, 6, 10, 9, 7, 6, 8, 6, 8], 1): ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = ws.cell(row=hr + 1, column=1)
ws2 = wb.create_sheet("Disagreements")
ws2.append(["Race", "PN", "Horse", "ML", "Champ P", "Champ odds", "Champ rk", "Tree P", "Tree odds", "Tree rk", "Δ pts", "Who likes"])
for c in ws2[1]: c.font = hdr
for _, h in big.iterrows():
    ws2.append([int(h["race"]), h["pn"], h["HorseName"], h["MornLineOddsifavailable"], h["p_cp"], h["odds_cp"], int(h["rank_p_cp"]), h["p_ch"], h["odds_ch"], int(h["rank_p_ch"]), h["dp"], "TREE" if h["dp"] > 0 else "CHAMP"])
    row = ws2[ws2.max_row]; row[4].number_format = "0.0%"; row[7].number_format = "0.0%"; row[5].number_format = "0.0"; row[8].number_format = "0.0"; row[10].number_format = "+0.0;-0.0"
for i, w in enumerate([6, 4, 22, 6, 9, 10, 8, 9, 9, 7, 8, 9], 1): ws2.column_dimensions[get_column_letter(i)].width = w
ws2.freeze_panes = "A2"
xout = HERE / "results" / f"daily_{a.track}_{d}.xlsx"; wb.save(xout); print("xlsx ->", xout.name)
