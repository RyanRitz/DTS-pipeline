"""For a pre-scored card (prescore_h2h.py), explain champion-vs-tree divergences at the variable level.
Champion: per-feature logit contributions via attribution._blend_contribution (exact scoring blend).
Tree:     per-feature logit contributions via LightGBM pred_contrib (TreeSHAP).
Both are made RELATIVE to the race mean, so they answer "why does this model rate this horse above/below the field".
Writes results/divergence_{TRK}_{date}.csv (all horses x all features, long) and prints the top drivers for
the biggest disagreements. Also builds results/prescore_{TRK}_{dates}.xlsx when --xlsx is given."""
import sys, argparse, warnings; warnings.simplefilter("ignore")
from pathlib import Path
import numpy as np, pandas as pd
HERE = Path(__file__).resolve().parent; sys.path.insert(0, str(HERE))
CH = HERE.parent / "Challenger" / "data"
import config, model_setup, score, model_registry as mr, attribution as at
from ingest_drf import load_drf; from features import engineer_features
import lightgbm as lgb

def contributions(track, date, drf_path):
    model_setup.setup_registry(config)
    cdir = Path(config.COEFF_DIR); cdir = cdir if cdir.exists() else HERE / "coefficients"
    mm, yyyy = date[4:6], date[:4]
    drf = load_drf(drf_path, track, mm + date[6:8], yyyy, validate=False)
    feats = engineer_features(drf.copy())
    sc = mr.get_scoring_models(track, config, date)
    cp = score.run_scoring(feats.copy(), cdir, sc)
    cp["pn"] = cp["ProgramNumberifavailable"].astype(str).str.strip(); cp["race"] = cp["Race"].astype(int)
    feats["pn"] = feats["ProgramNumberifavailable"].astype(str).str.strip(); feats["race"] = pd.to_numeric(feats["Race"], errors="coerce").astype(int)
    # --- champion contributions (same code path as the sheet's why-like) ---
    key = ["Track", "Date", "Race", "HorseName"]
    extra = [c for c in feats.columns if c not in cp.columns or c in key]
    merged = cp.merge(feats[extra], on=key, how="left")
    coeff_sets = at._load_coefficient_sets(sc, cdir, merged.columns)
    mplan = at._maiden_plan(sc); turf26 = at._is_oct26_turf(sc)
    rows = []
    for race, rg in merged.groupby("Race"):
        mid = rg["model"].iloc[0]; sub = coeff_sets.get(mid, {})
        if not sub: continue
        allf = sorted({f for c in sub.values() for f in c} & set(rg.columns))
        cr = {i: at._blend_contribution(rg.loc[i], mid, sub, allf, mplan, turf26) for i in rg.index}
        cdf = pd.DataFrame([cr[i] or {} for i in rg.index], index=rg.index).reindex(columns=allf).fillna(0.0)
        rel = cdf - cdf.mean()
        for i in rg.index:
            for f in allf:
                rows.append((int(race), rg.at[i, "pn"], f, float(cdf.at[i, f]), float(rel.at[i, f])))
    champ = pd.DataFrame(rows, columns=["race", "pn", "feature", "contrib_cp", "rel_cp"])
    # --- tree contributions (TreeSHAP, log-odds) ---
    models = {s: lgb.Booster(model_file=str(CH / f"sastree_{s}.txt")) for s in ["dirt", "turf", "maiden"]}
    feat = models["dirt"].feature_name()
    x = feats.copy()
    for v in feat: x[v] = pd.to_numeric(x[v], errors="coerce") if v in x.columns else np.nan
    S = x["Surface"].astype(str); RT = x["RaceType"].astype(str)
    x["seg"] = np.where(RT.isin(["M", "S"]), "maiden", np.where(S.isin(["T", "t"]), "turf", "dirt"))
    trows = []
    for s, m in models.items():
        k = x["seg"] == s
        if not k.any(): continue
        C = m.predict(x.loc[k, feat].values.astype(np.float32), pred_contrib=True)   # n x (F+1), last = bias
        cdf = pd.DataFrame(C[:, :-1], columns=feat, index=x.index[k])
        cdf["race"] = x.loc[k, "race"]; cdf["pn"] = x.loc[k, "pn"]
        for race, g in cdf.groupby("race"):
            rel = g[feat] - g[feat].mean()
            for i in g.index:
                for f in feat:
                    trows.append((int(race), g.at[i, "pn"], f, float(g.at[i, f]), float(rel.at[i, f])))
    tree = pd.DataFrame(trows, columns=["race", "pn", "feature", "contrib_ch", "rel_ch"])
    allc = champ.merge(tree, on=["race", "pn", "feature"], how="outer").fillna(0.0)
    # feature value for context
    vals = feats.melt(id_vars=["race", "pn"], value_vars=[f for f in feat if f in feats.columns], var_name="feature", value_name="value")
    allc = allc.merge(vals, on=["race", "pn", "feature"], how="left")
    allc["gap"] = allc["rel_ch"] - allc["rel_cp"]          # + = tree likes this feature more than champion does, relative to field
    return allc

def explain(pre, allc, top_n=8, min_gap_pct=6.0, max_horses=12):
    pre = pre.copy(); pre["dp"] = (pre["p_ch"] - pre["p_cp"]) * 100
    pre["pn"] = pre["pn"].astype(str)
    big = pre[pre["dp"].abs() >= min_gap_pct].copy(); big = big.loc[big["dp"].abs().sort_values(ascending=False).index].head(max_horses)
    lines = []
    for _, h in big.iterrows():
        a = allc[(allc["race"] == h["race"]) & (allc["pn"] == h["pn"])].copy()
        a = a.loc[a["gap"].abs().sort_values(ascending=False).index].head(top_n)
        lines.append(f"\nR{int(h['race'])} #{h['pn']} {h['HorseName']}  ML {h['MornLineOddsifavailable']}  champion {h['p_cp']*100:.1f}% (rk {int(h['rank_p_cp'])})  tree {h['p_ch']*100:.1f}% (rk {int(h['rank_p_ch'])})  Δ {h['dp']:+.1f} pts")
        lines.append(f"   {'feature':28s} {'value':>9s} {'champ rel':>10s} {'tree rel':>9s} {'gap':>7s}")
        for _, r in a.iterrows():
            v = r["value"]; vs = f"{v:9.3g}" if pd.notna(v) else "      nan"
            lines.append(f"   {r['feature']:28s} {vs} {r['rel_cp']:+10.3f} {r['rel_ch']:+9.3f} {r['gap']:+7.3f}")
    return "\n".join(lines), big

if __name__ == "__main__":
    ap = argparse.ArgumentParser(); ap.add_argument("--dates", required=True, help="comma list, e.g. 20261014,20261015")
    ap.add_argument("--track", default="KEE"); ap.add_argument("--xlsx", action="store_true"); a = ap.parse_args()
    sheets = {}; expl = {}
    for d in a.dates.split(","):
        drf = HERE.parent / a.track / "RAW_DATA" / "RACINGFORM" / d[:4] / f"{a.track}{d[4:]}.DRF"
        pre = pd.read_csv(HERE / "results" / f"prescore_{a.track}_{d}.csv", dtype={"pn": str})
        allc = contributions(a.track, d, str(drf))
        allc.to_csv(HERE / "results" / f"divergence_{a.track}_{d}.csv", index=False)
        txt, big = explain(pre, allc)
        open(HERE / "results" / f"divergence_{a.track}_{d}.txt", "w").write(f"{a.track} {d} — variable drivers of champion/tree divergence (relative-to-field logit contributions)\n" + txt)
        print(f"=== {d}"); print(txt)
        sheets[d] = pre; expl[d] = (txt, big, allc)
    if a.xlsx:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
        wb = Workbook(); wb.remove(wb.active)
        gold_cp = PatternFill("solid", fgColor="FFF2CC"); gold_ch = PatternFill("solid", fgColor="DDEBF7"); both = PatternFill("solid", fgColor="E2EFDA")
        hdr = Font(bold=True)
        for d, pre in sheets.items():
            ws = wb.create_sheet(f"{d[4:6]}-{d[6:]}")
            cols = ["race", "pn", "HorseName", "Surface", "RaceType", "MornLineOddsifavailable", "MornOddsAdj",
                    "p_cp", "odds_cp", "rank_p_cp", "gold_p_cp", "top50_p_cp",
                    "p_ch", "odds_ch", "rank_p_ch", "gold_p_ch", "top50_p_ch", "dp", "edge_p_cp", "edge_p_ch"]
            names = ["Race", "PN", "Horse", "Surf", "Type", "ML", "ML adj", "Champ P(win)", "Champ odds", "Champ rank", "Champ GOLD", "Champ top50",
                     "Tree P(win)", "Tree odds", "Tree rank", "Tree GOLD", "Tree top50", "Tree − Champ (pts)", "Champ edge vs ML", "Tree edge vs ML"]
            pre = pre.sort_values(["race", "p_cp"], ascending=[True, False]).copy(); pre["dp"] = (pre["p_ch"] - pre["p_cp"]) * 100
            ws.append(names)
            for c in ws[1]: c.font = hdr; c.alignment = Alignment(horizontal="center", wrap_text=True)
            for _, r in pre.iterrows():
                ws.append([r[c] if not isinstance(r[c], (bool, np.bool_)) else ("GOLD" if r[c] and "gold" in c else ("+" if r[c] else "")) for c in cols])
                row = ws.max_row
                if r["gold_p_cp"] and r["gold_p_ch"]: fill = both
                elif r["gold_p_cp"]: fill = gold_cp
                elif r["gold_p_ch"]: fill = gold_ch
                else: fill = None
                if fill:
                    for c in ws[row]: c.fill = fill
            for row in ws.iter_rows(min_row=2):
                row[7].number_format = "0.0%"; row[12].number_format = "0.0%"
                for j in (8, 13, 6): row[j].number_format = "0.0"
                for j in (17,): row[j].number_format = "+0.0;-0.0;0.0"
                for j in (18, 19): row[j].number_format = "0.00"
            widths = [5, 4, 22, 5, 5, 6, 7, 10, 9, 8, 9, 9, 10, 9, 8, 9, 9, 11, 10, 10]
            for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
            ws.freeze_panes = "D2"
            ws.cell(row=ws.max_row + 2, column=1, value="Yellow = champion gold · Blue = tree gold · Green = gold by both. Pre-scratch; gates shift when the field is set.").font = Font(italic=True)
        ws = wb.create_sheet("Divergence drivers")
        ws.append(["Card", "Race", "PN", "Horse", "ML", "Champ P", "Tree P", "Δ pts", "Feature", "Value", "Champ rel contrib", "Tree rel contrib", "Gap (tree − champ)"])
        for c in ws[1]: c.font = hdr
        for d, (txt, big, allc) in expl.items():
            for _, h in big.iterrows():
                a_ = allc[(allc["race"] == h["race"]) & (allc["pn"] == h["pn"])]
                a_ = a_.loc[a_["gap"].abs().sort_values(ascending=False).index].head(8)
                for _, r in a_.iterrows():
                    ws.append([f"{d[4:6]}-{d[6:]}", int(h["race"]), h["pn"], h["HorseName"], h["MornLineOddsifavailable"], h["p_cp"], h["p_ch"], (h["p_ch"] - h["p_cp"]) * 100,
                               r["feature"], r["value"] if pd.notna(r["value"]) else None, r["rel_cp"], r["rel_ch"], r["gap"]])
                ws.append([])
        for row in ws.iter_rows(min_row=2):
            if row[5].value is not None:
                row[5].number_format = "0.0%"; row[6].number_format = "0.0%"; row[7].number_format = "+0.0;-0.0"
                for j in (9, 10, 11, 12): row[j].number_format = "0.000"
        for i, w in enumerate([7, 5, 4, 22, 6, 8, 8, 7, 26, 9, 14, 14, 14], 1): ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        ws.cell(row=ws.max_row + 2, column=1, value="Contributions are log-odds, relative to the race mean (+ = this model rates the horse above the field on that variable). Champion = coefficient × value through the exact scoring blend; Tree = TreeSHAP. Gap = where the two models read the same variable differently.").font = Font(italic=True)
        out = HERE / "results" / f"prescore_{a.track}_{a.dates.replace(',', '_')}.xlsx"
        wb.save(out); print("xlsx ->", out)
